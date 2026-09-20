"""Turn an uploaded file into plain text for Qwen.

Detection trusts file CONTENT over the file name, because xlsx and docx are both
ZIP files ("PK...") and used to be confused, which crashed python-docx.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import logging
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from config import (
    ENABLE_OCR,
    MAX_PDF_PAGES,
    MAX_SHEET_ROWS,
    PDF_FORCE_OCR,
    PDF_MIN_CHARS_PER_PAGE,
    PDF_OCR_DPI,
)

log = logging.getLogger("extractors")

SUPPORTED = "PDF, Word (DOCX/DOC), Excel (XLSX/XLS), CSV/TSV, TXT/MD, JSON"


class UnsupportedFile(ValueError):
    pass


@dataclass
class Extracted:
    text: str
    kind: str
    method: str
    warnings: list[str] = field(default_factory=list)


# ================================================================== helpers
_BIDI_CONTROLS = dict.fromkeys(map(ord, "\u200e\u200f\u202a\u202b\u202c\u202d\u202e\u2066\u2067\u2068\u2069\ufeff"), None)


def clean_text(text):
    """Normalize a block of text while keeping line breaks."""
    if not text:
        return ""
    # NFKC turns Arabic presentation forms (common in PDFs) into normal letters.
    text = unicodedata.normalize("NFKC", str(text)).translate(_BIDI_CONTROLS)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    lines = [" ".join(line.split()) for line in text.split("\n")]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()




def _clean_docx_paragraph_text(text):
    """Normalize DOCX paragraph text while preserving obvious visual columns.

    Word reports often place multiple headings on one paragraph using a long run of
    spaces.  The generic clean_text() intentionally collapses whitespace, which
    destroys that structure.  Convert only *large* horizontal gaps/tabs to a
    neutral column separator before normalizing the remaining whitespace.
    """
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", str(text)).translate(_BIDI_CONTROLS)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    # 4+ spaces or one/more tabs are layout, not ordinary word spacing.
    text = re.sub(r"(?:\t+| {4,})", " | ", text)
    lines = [" ".join(line.split()) for line in text.split("\n")]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()

def _cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.6g}"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return " ".join(str(value).split())


def _row_line(values):
    cells = [_cell(v) for v in values]
    while cells and not cells[-1]:
        cells.pop()
    # Never drop cells by text: two columns can legitimately hold the same value
    # (e.g. 3000 | 3000) and dropping one shifts the row under the wrong headers.
    return " | ".join(cells) if any(cells) else ""


def decode_bytes(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for enc in ("utf-8-sig", "cp1256"):  # cp1256 = Arabic Windows exports
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


# ================================================================== detection
EXT_KIND = {
    ".pdf": "pdf", ".docx": "docx", ".docm": "docx", ".doc": "doc",
    ".xlsx": "xlsx", ".xlsm": "xlsx", ".xls": "xls",
    ".csv": "csv", ".tsv": "tsv", ".txt": "text", ".md": "text", ".json": "json",
    ".pptx": "pptx", ".ppt": "ppt",
}
MIME_KIND = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/msword": "doc",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel": "xls",
    "text/csv": "csv", "text/tab-separated-values": "tsv",
    "text/plain": "text", "text/markdown": "text",
    "application/json": "json", "text/json": "json",
}


def _sniff(path: Path):
    with open(path, "rb") as f:
        head = f.read(8192)
    if b"%PDF" in head[:1024]:
        return "pdf"
    if head.startswith(b"PK"):
        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
        except zipfile.BadZipFile:
            return None
        if any(n.startswith("word/") for n in names):
            return "docx"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        if any(n.startswith("ppt/") for n in names):
            return "pptx"
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "ole"  # legacy .doc or .xls
    if head.startswith((b"\xff\xfe", b"\xfe\xff")) or b"\x00" not in head:
        stripped = decode_bytes(head).lstrip()
        return "json" if stripped[:1] in ("{", "[") else "text"
    return None


def detect_kind(path: Path, name=None, mime=None):
    sniffed = _sniff(path)
    by_ext = EXT_KIND.get(Path(name or path.name).suffix.lower())
    by_mime = MIME_KIND.get((mime or "").split(";")[0].strip().lower())
    hint = by_ext or by_mime

    if sniffed in {"pdf", "docx", "xlsx", "pptx", "zip"}:
        return sniffed
    if sniffed == "ole":
        return hint if hint in {"doc", "xls"} else "ole"
    if sniffed in {"text", "json"}:
        return hint if hint in {"csv", "tsv", "text", "json"} else sniffed
    return hint or "unknown"


# ================================================================== extractors
def _extract_text_file(path, kind):
    text = decode_bytes(path.read_bytes())
    if kind == "json":
        try:
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass  # still useful as plain text
    if kind in {"csv", "tsv"}:
        delimiter = "\t" if kind == "tsv" else None
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        lines = []
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
            if i >= MAX_SHEET_ROWS:
                lines.append(f"... (truncated after {MAX_SHEET_ROWS} rows)")
                break
            line = _row_line(row)
            if line:
                lines.append(line)
        text = "\n".join(lines)
    return clean_text(text)


def _extract_docx(path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    chunks, seen = [], set()

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = _clean_docx_paragraph_text(Paragraph(child, doc).text)
            if text:
                chunks.append(text)
                seen.add(text)
        elif child.tag == qn("w:tbl"):
            rows = []
            for row in Table(child, doc).rows:
                cells, seen_tc = [], set()
                for cell in row.cells:  # a horizontally merged cell is the SAME <w:tc> repeated
                    if id(cell._tc) in seen_tc:
                        continue
                    seen_tc.add(id(cell._tc))
                    cells.append(clean_text(cell.text).replace("\n", " "))
                line = _row_line(cells)
                if line:
                    rows.append(line)
            if rows:
                chunks.extend(["===== TABLE START =====", *rows, "===== TABLE END ====="])
                seen.update(rows)

    # Text boxes / shapes are invisible to paragraph.text but common in Arabic reports.
    box_lines = []
    for p in doc.element.body.iter(qn("w:txbxContent")):
        text = _clean_docx_paragraph_text("".join(t.text or "" for t in p.iter(qn("w:t"))))
        if text and text not in seen:
            seen.add(text)
            box_lines.append(text)
    if box_lines:
        chunks.extend(["===== TEXT BOXES =====", *box_lines])

    return "\n".join(chunks).strip()


def _extract_doc(path):
    if shutil.which("antiword"):
        res = subprocess.run(["antiword", "-w", "0", "-m", "UTF-8.txt", str(path)],
                             capture_output=True, timeout=60)
        if res.returncode == 0 and res.stdout.strip():
            return clean_text(res.stdout.decode("utf-8", errors="replace")), "antiword"
    if shutil.which("soffice"):
        with tempfile.TemporaryDirectory() as out:
            subprocess.run(["soffice", "--headless", "--convert-to", "docx", "--outdir", out, str(path)],
                           capture_output=True, timeout=120)
            converted = next(Path(out).glob("*.docx"), None)
            if converted:
                return _extract_docx(converted), "soffice+python-docx"
    raise UnsupportedFile("Legacy .doc could not be read (antiword/LibreOffice not available). Save it as .docx.")


def _extract_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    chunks = []
    try:
        for ws in wb.worksheets:
            rows = []
            for i, values in enumerate(ws.iter_rows(values_only=True)):
                if len(rows) >= MAX_SHEET_ROWS:
                    rows.append(f"... (truncated after {MAX_SHEET_ROWS} rows)")
                    break
                line = _row_line(values)
                if line:
                    rows.append(clean_text(line))
            if rows:
                chunks.extend([f"===== SHEET: {ws.title} =====", "===== TABLE START =====", *rows, "===== TABLE END ====="])
    finally:
        wb.close()
    return "\n".join(chunks).strip()


def _extract_xls(path):
    import xlrd

    book = xlrd.open_workbook(str(path))
    chunks = []
    for sheet in book.sheets():
        rows = []
        for r in range(min(sheet.nrows, MAX_SHEET_ROWS)):
            values = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        values.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                        continue
                    except Exception:
                        pass
                values.append(cell.value)
            line = _row_line(values)
            if line:
                rows.append(clean_text(line))
        if rows:
            chunks.extend([f"===== SHEET: {sheet.name} =====", "===== TABLE START =====", *rows, "===== TABLE END ====="])
    return "\n".join(chunks).strip()


# ------------------------------------------------------------------ PDF
def _letters(text):
    return sum(ch.isalpha() for ch in text)


def _arabic_looks_reversed(text):
    """Some PDFs store Arabic in visual order: 'البرنامج' comes out as 'جمانربلا'."""
    words = re.findall(r"[\u0621-\u064A]{3,}", text)
    if len(words) < 20:
        return False
    starts = sum(w.startswith("ال") for w in words)
    ends = sum(w.endswith("لا") for w in words)
    return ends > max(5, starts * 1.5)


def _page_needs_ocr(text):
    if _letters(text) < PDF_MIN_CHARS_PER_PAGE:
        return True
    if text.count("\ufffd") > 0.02 * len(text):
        return True
    return _arabic_looks_reversed(text)


def _extract_pdf(path):
    import pymupdf

    warnings = []
    with pymupdf.open(str(path)) as doc:
        if doc.needs_pass:
            raise UnsupportedFile("PDF is password-protected.")
        total = doc.page_count
        if total > MAX_PDF_PAGES:
            warnings.append(f"PDF has {total} pages; only the first {MAX_PDF_PAGES} were read.")
        page_texts = [clean_text(doc[i].get_text("text", sort=True)) for i in range(min(total, MAX_PDF_PAGES))]

        ocr_pages = [i for i, t in enumerate(page_texts) if PDF_FORCE_OCR or _page_needs_ocr(t)]
        method = "pdf-text-layer"

        if ocr_pages and ENABLE_OCR:
            from pipeline.ocr import ocr_image

            with tempfile.TemporaryDirectory() as tmp:
                for i in ocr_pages:
                    png = Path(tmp) / f"page_{i + 1}.png"
                    doc[i].get_pixmap(dpi=PDF_OCR_DPI).save(str(png))
                    try:
                        ocr_text = clean_text(ocr_image(png))
                    except Exception as exc:  # keep the text layer instead of failing the file
                        log.exception("OCR failed on page %s", i + 1)
                        warnings.append(f"OCR failed on page {i + 1}: {exc}")
                        continue
                    if _letters(ocr_text) >= _letters(page_texts[i]) or _arabic_looks_reversed(page_texts[i]):
                        page_texts[i] = ocr_text
            method = "tesseract-ocr" if len(ocr_pages) == len(page_texts) else "pdf-text-layer+tesseract-ocr"
        elif ocr_pages:
            warnings.append(f"{len(ocr_pages)} page(s) look scanned but OCR is disabled.")

    pages = [f"===== PAGE {i} =====\n{t}" for i, t in enumerate(page_texts, 1) if t]
    return "\n\n".join(pages).strip(), method, warnings


# ================================================================== entry point
def extract_file(path: Path, name=None, mime=None) -> Extracted:
    path = Path(path)
    kind = detect_kind(path, name, mime)
    log.info("file=%r kind=%s size=%d", name or path.name, kind, path.stat().st_size)

    if kind == "pdf":
        text, method, warnings = _extract_pdf(path)
        return Extracted(text, kind, method, warnings)
    if kind == "docx":
        return Extracted(_extract_docx(path), kind, "python-docx")
    if kind == "doc":
        text, method = _extract_doc(path)
        return Extracted(text, kind, method)
    if kind == "xlsx":
        return Extracted(_extract_xlsx(path), kind, "openpyxl")
    if kind == "xls":
        return Extracted(_extract_xls(path), kind, "xlrd")
    if kind == "ole":  # legacy Office file without a usable name
        try:
            return Extracted(_extract_xls(path), "xls", "xlrd")
        except Exception:
            text, method = _extract_doc(path)
            return Extracted(text, "doc", method)
    if kind in {"csv", "tsv", "text", "json"}:
        return Extracted(_extract_text_file(path, kind), kind, "text")
    if kind in {"pptx", "ppt"}:
        raise UnsupportedFile(f"PowerPoint files are not supported. Supported: {SUPPORTED}.")
    raise UnsupportedFile(f"Unsupported file type ({name or path.name}). Supported: {SUPPORTED}.")

"""Offline tests: no GPU, no Qwen, no live OCR engine required.

Run:  python -m pytest -q tests        (or)   python tests/test_offline.py
"""
import base64
import json
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["PRELOAD_MODELS"] = "false"
os.environ["ENABLE_OCR"] = "false"

import pipeline.pipeline as pl  # noqa: E402
from pipeline.extractors import detect_kind, extract_file  # noqa: E402
from pipeline.inputs import InputError, normalize_request  # noqa: E402
from pipeline.json_utils import parse_program_items  # noqa: E402
from pipeline.validator import FIELDS, clean_program  # noqa: E402

TMP = Path(tempfile.mkdtemp())


# ------------------------------------------------------------------ fixtures
def make_docx():
    from docx import Document
    from docx.oxml import parse_xml

    doc = Document()
    doc.add_paragraph("تقرير جمعية الخير السنوي 2025")
    table = doc.add_table(rows=2, cols=3)
    for i, v in enumerate(["اسم المشروع", "المستفيدون", "الميزانية"]):
        table.rows[0].cells[i].text = v
    for i, v in enumerate(["سقيا الماء", "1200", "45000"]):
        table.rows[1].cells[i].text = v
    textbox = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape><v:textbox><w:txbxContent>'
        "<w:p><w:r><w:t>نص داخل مربع نص</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    doc.element.body.append(parse_xml(textbox))
    path = TMP / "report.docx"
    doc.save(path)
    return path


def make_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "المشاريع"
    ws.append(["اسم المشروع", "السنة", "الميزانية"])
    ws.append(["كفالة يتيم", 2024, 150000.0])
    ws.append([None, None, None])
    path = TMP / "budget.xlsx"
    wb.save(path)
    return path


def make_pdf():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Clean Water Initiative 2025. Beneficiaries: 1200 households. "
                               "Budget 45000 USD for well construction.", fontsize=11)
    path = TMP / "brief.pdf"
    doc.save(path)
    return path


def b64(path):
    return base64.b64encode(Path(path).read_bytes()).decode()


# ------------------------------------------------------------------ detection & extraction
def test_xlsx_is_not_mistaken_for_docx():
    xlsx = make_xlsx()
    assert detect_kind(xlsx, "budget.xlsx") == "xlsx"
    assert detect_kind(xlsx, "wrong-name.docx") == "xlsx"  # content wins
    assert detect_kind(xlsx, None) == "xlsx"


def test_docx_tables_and_textboxes():
    r = extract_file(make_docx(), "report.docx")
    assert r.kind == "docx"
    assert "سقيا الماء | 1200 | 45000" in r.text
    assert "TABLE START" in r.text
    assert "نص داخل مربع نص" in r.text


def test_xlsx_values():
    r = extract_file(make_xlsx(), "budget.xlsx")
    assert "كفالة يتيم | 2024 | 150000" in r.text, r.text
    assert "SHEET: المشاريع" in r.text


def test_xls_legacy():
    import pytest
    xlwt = pytest.importorskip("xlwt")  # optional, only to build the fixture
    wb = xlwt.Workbook()
    ws = wb.add_sheet("s")
    ws.write(0, 0, "مشروع")
    ws.write(0, 1, 300)
    path = TMP / "old.xls"
    wb.save(str(path))
    r = extract_file(path, None)  # no name -> OLE sniff -> xls
    assert r.kind == "xls" and "مشروع | 300" in r.text


def test_pdf_text_layer():
    r = extract_file(make_pdf(), "brief.pdf")
    assert r.kind == "pdf" and r.method == "pdf-text-layer"
    assert "Clean Water Initiative" in r.text


def test_csv_cp1256_and_txt():
    csv_path = TMP / "a.csv"
    csv_path.write_bytes("الاسم;العدد\nتدريب;50\n".encode("cp1256"))
    r = extract_file(csv_path, "a.csv")
    assert "تدريب | 50" in r.text, r.text
    txt = TMP / "a.txt"
    txt.write_text("برنامج محو الأمية", encoding="utf-8")
    assert extract_file(txt, "a.txt").text == "برنامج محو الأمية"


def test_arabic_presentation_forms_normalized():
    from pipeline.extractors import clean_text
    assert clean_text("\ufe8d\ufedf\ufe92\ufeae\ufee7\ufe8e\ufee3\ufe9e") == "البرنامج"


def test_reversed_arabic_detection():
    from pipeline.extractors import _arabic_looks_reversed
    normal = " ".join(["البرنامج", "المستفيدين", "الجمعية", "التعليم"] * 6)
    reversed_ = " ".join(w[::-1] for w in normal.split())
    assert not _arabic_looks_reversed(normal)
    assert _arabic_looks_reversed(reversed_)


# ------------------------------------------------------------------ input shapes
def test_contract_input_string_with_attachments():
    req = normalize_request({
        "run_id": "42", "type": "program_autofill",
        "input": json.dumps({"text": "انظر الملف المرفق", "files": ["brief.pdf"]}),
        "attachments": [b64(make_pdf())],
    })
    assert req.run_id == "42" and req.text == "انظر الملف المرفق"
    assert len(req.files) == 1 and req.files[0].name == "brief.pdf" and req.files[0].data


def test_contract_outer_files_used_as_content():
    req = normalize_request({
        "type": "program_autofill",
        "input": '{"files":["a.xlsx","a.xlsx"]}',
        "files": [{"data": b64(make_xlsx())}, {"data": b64(make_xlsx())}],
    })
    assert [f.name for f in req.files] == ["a.xlsx", "a.xlsx"]  # duplicate names allowed


def test_recommended_shape_and_legacy():
    req = normalize_request({"text": "x", "files": [{"name": "b.docx", "data": "data:application/octet-stream;base64," + b64(make_docx())}]})
    assert req.type == "program_extract" and req.files[0].data
    legacy = normalize_request({"file_url": "https://example.com/r.pdf"})
    assert legacy.files[0].url == "https://example.com/r.pdf"


def test_empty_input_and_names_only_are_input_errors():
    for bad in ({"input": "{}"}, {"input": '{"files":["brief.pdf"]}'}, None, "not json"):
        try:
            job = normalize_request(bad)
            pl.collect_sources(job)
        except InputError:
            continue
        raise AssertionError(f"expected InputError for {bad!r}")


# ------------------------------------------------------------------ model output handling
def test_parse_fenced_and_truncated_json():
    fenced = '```json\n{"programs":[{"name":"أ","type":"program"},]}\n```'
    assert parse_program_items(fenced)[0]["name"] == "أ"
    truncated = '{"programs":[{"name":"أ"},{"name":"ب","year":2025},{"name":"ج","descr'
    items = parse_program_items(truncated, truncated=True)
    assert [i["name"] for i in items] == ["أ", "ب"]
    assert parse_program_items('<think>\n</think>\n{"name":"single"}')[0]["name"] == "single"


def test_coercion():
    p = clean_program({
        "type": "مشروع", "name": "  سقيا   الماء ", "year": "٢٠٢٥م",
        "beneficiaries_count": "1,200 مستفيد", "budget": "٤٥٬٠٠٠٫٥ ريال",
        "description": "", "notes": "null", "target_audience": ["أسر", "أيتام"],
    })
    assert p == {
        "type": "project", "name": "سقيا الماء", "year": 2025, "description": None,
        "beneficiaries_count": 1200, "budget": 45000.5, "beneficiary_value": None,
        "target_audience": "أسر، أيتام", "delivery_method": None, "notes": None,
    }
    assert clean_program({"name": "x", "year": 1446, "budget": "1.5 مليون"})["budget"] == 1500000
    assert clean_program({"name": "x", "year": 1446})["year"] is None  # Hijri year is out of range
    assert clean_program({"name": ""}) is None


# ------------------------------------------------------------------ end-to-end with a fake model
def _fake_model(response_items):
    def fake(document, max_new_tokens, chunk_chars):
        fake.document = document
        return pl.merge_programs(pl.clean_programs(response_items)), {"truncated": False, "chunk_errors": []}
    return fake


def test_autofill_end_to_end_contract_shape():
    pl._run_model = _fake_model([
        {"name": "مشروع صغير", "type": "project"},
        {"name": "Clean Water Initiative", "type": "program", "year": "2025", "budget": "45,000", "beneficiaries_count": 1200},
    ])
    out = pl.process_request({
        "run_id": "42", "type": "program_autofill",
        "input": json.dumps({"text": "انظر الملف المرفق", "files": ["brief.pdf", "budget.xlsx"]}),
        "attachments": [b64(make_pdf()), b64(make_xlsx())],
    })
    assert list(out) == list(FIELDS), out.keys()
    assert out["name"] == "Clean Water Initiative" and out["year"] == 2025 and out["budget"] == 45000.0
    assert "تم العثور على 2" in out["notes"]
    doc = pl._run_model.document
    assert "USER TEXT" in doc and "FILE 1: brief.pdf" in doc and "FILE 2: budget.xlsx" in doc


def test_autofill_no_programs_does_not_invent_fallback():
    pl._run_model = _fake_model([])
    out = pl.process_request({"type": "program_autofill", "files": [{"name": "كفالة الأيتام.docx", "data": b64(make_docx())}]})
    assert out["name"] is None and out["type"] is None and out["notes"]


def test_bad_file_is_warning_not_crash():
    pl._run_model = _fake_model([{"name": "برنامج"}])
    pptx_like = TMP / "deck.pptx"
    import zipfile
    with zipfile.ZipFile(pptx_like, "w") as z:
        z.writestr("ppt/presentation.xml", "<x/>")
    out = pl.process_request({"type": "program_autofill", "text": "برنامج", "files": [{"name": "deck.pptx", "data": b64(pptx_like)}]})
    assert out["name"] == "برنامج" and "deck.pptx" in out["notes"]


def test_extract_text_type():
    out = pl.process_request({"type": "extract_text", "files": [{"name": "budget.xlsx", "data": b64(make_xlsx())}]})
    assert "كفالة يتيم" in out["text"] and out["sources"][0]["kind"] == "xlsx"


def test_scanned_pdf_goes_to_ocr():
    import types
    import pymupdf
    import pipeline.extractors as ex

    doc = pymupdf.open()
    doc.new_page()  # blank page = no text layer, like a scan
    path = TMP / "scan.pdf"
    doc.save(path)

    calls = []
    fake = types.ModuleType("pipeline.ocr")
    fake.ocr_image = lambda png: calls.append(Path(png).exists()) or "برنامج من OCR"
    sys.modules["pipeline.ocr"] = fake
    ex.ENABLE_OCR = True
    try:
        r = extract_file(path, "scan.pdf")
    finally:
        ex.ENABLE_OCR = False
        del sys.modules["pipeline.ocr"]
    assert calls == [True] and r.method == "tesseract-ocr" and "برنامج من OCR" in r.text


def test_equal_adjacent_cells_do_not_shift_columns():
    from pipeline.extractors import _row_line
    assert _row_line(["كفالة الأيتام", "3000", "3000", "29991"]) == "كفالة الأيتام | 3000 | 3000 | 29991"


def test_docx_merged_cells_not_repeated_but_equal_values_kept():
    from docx import Document
    doc = Document()
    t = doc.add_table(rows=2, cols=4)
    a = t.rows[0].cells[0].merge(t.rows[0].cells[1]); a.text = "عنوان مدمج"
    for i, v in enumerate(["إفطار صائم", "3000", "3000", "29991"]):
        t.rows[1].cells[i].text = v
    path = TMP / "merged.docx"; doc.save(path)
    text = extract_file(path, "merged.docx").text
    assert "عنوان مدمج\n" in text and "عنوان مدمج | عنوان مدمج" not in text, text
    assert "إفطار صائم | 3000 | 3000 | 29991" in text


def test_multiple_top_level_objects():
    out = '{"name":"أ"}\n{"name":"ب"},\n{"programs":[{"name":"ج"}]}'
    assert [i["name"] for i in parse_program_items(out)] == ["أ", "ب", "ج"]


def test_program_extract_is_default_and_returns_ordered_list():
    pl._run_model = _fake_model([
        {"name": "مشروع كفالة الأيتام", "type": "project", "budget": "150,000"},
        {"name": "Clean Water Initiative", "type": "program", "year": 2025, "budget": 45000},
        {"name": "برنامج محو الأمية", "type": "برنامج"},
    ])
    out = pl.process_request({"text": "تقرير"})
    assert isinstance(out, list) and len(out) == 3
    assert [o["type"] for o in out] == ["program", "program", "project"]
    assert all(list(o) == list(FIELDS) for o in out)
    assert out[0]["budget"] == 45000.0 and isinstance(out[0]["budget"], float)
    assert json.dumps(out, ensure_ascii=False)  # serializable


def test_grounding_drops_memorized_training_item():
    from pipeline.grounding import ground_programs
    bad = [{
        "type": "project", "name": "مشروع كفالة الأيتام", "year": 2025,
        "description": "إفطار الصائمين بالشراكة مع مؤسسة سليمان الراجحي الخيرية",
        "beneficiaries_count": 3000, "budget": 29991,
        "beneficiary_value": "إفطار الصائمين", "target_audience": "صائم",
        "delivery_method": "إفطار الصائمين", "notes": None,
    }]
    doc = "اسم المشروع | عدد المستفيدين | الميزانية\nإفطار الصائمين | 3000 | 29991"
    out, meta = ground_programs(pl.clean_programs(bad), doc)
    assert out == [] and meta["grounded_dropped"] == 1


def test_grounding_nulls_fields_not_on_source_row():
    from pipeline.grounding import ground_programs
    candidate = [{
        "type": "project", "name": "إفطار الصائمين", "year": 2025,
        "description": "وصف مخترع", "beneficiaries_count": 3000, "budget": 29991,
        "target_audience": "صائم",
    }]
    doc = "اسم المشروع | عدد المستفيدين | الميزانية | السنة\nإفطار الصائمين | 3000 | 29991 | 2025"
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    assert len(out) == 1
    assert out[0]["description"] is None and out[0]["target_audience"] is None
    assert out[0]["beneficiaries_count"] == 3000 and out[0]["budget"] == 29991.0


if __name__ == "__main__":
    failures = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS {name}")
            except ModuleNotFoundError as exc:
                print(f"SKIP {name}: {exc}")
            except Exception as exc:
                failures += 1
                print(f"FAIL {name}: {type(exc).__name__}: {exc}")
    sys.exit(1 if failures else 0)


def test_grounding_prefers_detail_occurrence_over_toc_and_keeps_context_fields():
    from pipeline.grounding import ground_programs
    doc = """16 1.مبادرة إفطار الصائمين
17 2.مبادرة سقيا الحاج والمعتمر
التقرير نصف السنوي لعام 2026
مبادرة إفطار الصائمين
تقديم وجبات جافة للصائمين خلال شهر رمضان المبارك
عدد المستفيدين
415900
الفئة المستهدفة
ضيوف الرحمن من الرجال والنساء
مكان التنفيذ
مصليات الحرم المكي الشريف
مبادرة سقيا الحاج والمعتمر
توزيع عبوات الماء البارد
عدد المستفيدين
534640
"""
    candidate = [{
        "type": "project", "name": "مبادرة إفطار الصائمين", "year": 2026,
        "description": "تقديم وجبات جافة للصائمين خلال شهر رمضان المبارك",
        "beneficiaries_count": 415900, "budget": None,
        "beneficiary_value": None,
        "target_audience": "ضيوف الرحمن من الرجال والنساء",
        "delivery_method": "مصليات الحرم المكي الشريف", "notes": None,
    }]
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    # v8.4 recovers other explicit entities too; it no longer throws away a
    # source-backed item merely because the model omitted it.
    assert [x["name"] for x in out] == ["مبادرة إفطار الصائمين", "مبادرة سقيا الحاج والمعتمر"]
    first, second = out
    assert first["beneficiaries_count"] == 415900
    assert first["description"] == "تقديم وجبات جافة للصائمين خلال شهر رمضان المبارك"
    assert first["target_audience"] == "ضيوف الرحمن من الرجال والنساء"
    assert first["delivery_method"] == "مصليات الحرم المكي الشريف"
    assert first["year"] == 2026
    assert second["beneficiaries_count"] == 534640
    assert second["delivery_method"] is None  # no leakage from the previous item


def test_grounding_does_not_accept_toc_sequence_as_beneficiary_count():
    from pipeline.grounding import ground_programs
    doc = """36 3.الدليل الإرشادي للجان الثقافية
الدليل الإرشادي للجان الثقافية
دليل إرشادي شامل يضم توجيهات تنظيمية
عدد المستفيدين
34 حملة مستفيدة
مكان التنفيذ
عن بعد
"""
    candidate = [{
        "type": "project", "name": "الدليل الإرشادي للجان الثقافية",
        "beneficiaries_count": 3,
    }]
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    assert len(out) == 1
    assert out[0]["beneficiaries_count"] == 34


def test_v6_structural_recovery_builds_programs_and_projects_from_document():
    from pipeline.grounding import ground_programs
    doc = """14 البرامج والمبادرات
15 أولا: برامج الإطعام
16 1.مبادرة إفطار الصائمين
17 2.مبادرة سقيا الماء
18 3.دورة تدريب المتطوعات

أولًا: برامج الإطعام
وصف برنامج الإطعام من المصدر.
مبادرة إفطار الصائمين
تقديم وجبات جافة للصائمين
عدد المستفيدين
415900
التقرير
ف 16
مبادرة سقيا الماء
توزيع عبوات الماء
عدد المستفيدين
534640
التقرير
ف 17
ثانيًا: برامج التدريب
تأهيل المتطوعات لخدمة المستفيدين.
دورة تدريب المتطوعات
دورة تدريبية عملية
عدد المستفيدين
67
التقرير
ف 18
تكريم المتطوعات
فعالية جانبية ليست ضمن قائمة البرامج والمشاريع.
"""
    out, meta = ground_programs([], doc)
    assert meta["structured_recovery"] is True
    assert [x["name"] for x in out if x["type"] == "program"] == ["برامج الإطعام", "برامج التدريب"]
    projects = [x for x in out if x["type"] == "project"]
    assert [x["name"] for x in projects] == [
        "مبادرة إفطار الصائمين", "مبادرة سقيا الماء", "دورة تدريب المتطوعات"
    ]
    assert [x["beneficiaries_count"] for x in projects] == [415900, 534640, 67]
    assert projects[0]["description"] == "تقديم وجبات جافة للصائمين"
    assert all("تكريم المتطوعات" != x["name"] for x in out)


def test_v7_multicolumn_rtl_reconstructs_target_and_delivery():
    from pipeline.grounding import _infer_multicolumn_target_delivery

    cases = [
        (
            """عدد المتطوعين ساعات التطوع مكان التنفيذ الفئة المستهدفة
مقر الجمعية ، ضيوف الرحمن من
12 116 الفنادق الرجال والنساء.
""",
            "ضيوف الرحمن من الرجال والنساء.",
            "مقر الجمعية الفنادق",
        ),
        (
            """عدد المتطوعين ساعات التطوع مكان التنفيذ الفئة المستهدفة
مقر الجمعية ، ضيوف الرحمن من
44 206 الفنادق ،الحملات النساء
""",
            "ضيوف الرحمن من النساء",
            "مقر الجمعية الفنادق ،الحملات",
        ),
        (
            """عدد المتطوعين ساعات التطوع مكان التنفيذ الفئة المستهدفة
مستشفى النور المر يضات من ضيفات
3 6 التخصصي بمكة الرحمن
المكرمة
""",
            "المر يضات من ضيفات الرحمن",
            "مستشفى النور التخصصي بمكة المكرمة",
        ),
        (
            """عدد المتطوعين ساعات التطوع مكان التنفيذ الفئة المستهدفة
مقر الجمعية ، أطفال زائرات البيت
9 58 الفنادق الحرام
""",
            "أطفال زائرات البيت الحرام",
            "مقر الجمعية الفنادق",
        ),
        (
            """عدد المتطوعين ساعات التطوع مكان التنفيذ الفئة المستهدفة
العاملون والمهتمون
2 12 غرفة مكة المكرمة
بالقطاع غير الربحي
""",
            "العاملون والمهتمون بالقطاع غير الربحي",
            "غرفة مكة المكرمة",
        ),
        (
            """الفئة المستهدفة مكان التنفيذ
ضيوف الرحمن من الرجال والنساء مصليات الحرم المكي الشريف
""",
            "ضيوف الرحمن من الرجال والنساء",
            "مصليات الحرم المكي الشريف",
        ),
        (
            """عدد المتطوعين ساعات التطوع مكان التنفيذ الفئة المستهدفة
مقر الجمعية ، الحاجات والمعتمرات
90 598 الفنادق ،الحملات والزائرات
""",
            "الحاجات والمعتمرات والزائرات",
            "مقر الجمعية الفنادق ،الحملات",
        ),
        (
            """عدد المتطوعين ساعات التطوع مكان التنفيذ الفئة المستهدفة
مقر الجمعية ،حملات حملات حجاج الداخل
5 100 الحج والخار ج من العرب
""",
            "حملات حجاج الداخل والخار ج من العرب",
            "مقر الجمعية ،حملات الحج",
        ),
    ]
    for context, target, delivery in cases:
        got_target, got_delivery = _infer_multicolumn_target_delivery(context)
        assert got_target == target, (got_target, target, context)
        assert got_delivery == delivery, (got_delivery, delivery, context)


def test_v7_beneficiary_value_never_duplicates_target_audience():
    from pipeline.grounding import ground_programs
    doc = """مبادرة إفطار الصائمين
تقديم وجبات جافة للصائمين
عدد المستفيدين
415900
الفئة المستهدفة مكان التنفيذ
ضيوف الرحمن من الرجال والنساء مصليات الحرم المكي الشريف
"""
    candidate = [{
        "type": "project", "name": "مبادرة إفطار الصائمين", "year": 2026,
        "beneficiary_value": "ضيوف الرحمن من الرجال والنساء",
        "target_audience": "ضيوف الرحمن من الرجال والنساء",
    }]
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    assert len(out) == 1
    assert out[0]["beneficiary_value"] is None


def test_v7_keeps_real_beneficiary_value_when_grounded_and_not_audience():
    from pipeline.grounding import ground_programs
    doc = """مبادرة إفطار الصائمين
تقديم وجبات جافة للصائمين
عدد المستفيدين
415900
الفئة المستهدفة
ضيوف الرحمن
"""
    candidate = [{
        "type": "project", "name": "مبادرة إفطار الصائمين",
        "beneficiary_value": "وجبات جافة", "target_audience": "ضيوف الرحمن",
    }]
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    assert len(out) == 1
    assert out[0]["beneficiary_value"] == "وجبات جافة"


def test_v82_notes_require_explicit_source_label():
    from pipeline.grounding import ground_programs
    doc = """دورة إدارة استراتيجية لحملة الحج
دورة تدريبية مقدمة لمشرفات الحج
الفئة المستهدفة مكان التنفيذ
مشرفات اللجان الثقافية عن بعد
"""
    candidate = [{
        "type": "project", "name": "دورة إدارة استراتيجية لحملة الحج",
        "notes": "عن بعد", "target_audience": "مشرفات اللجان الثقافية",
        "delivery_method": "عن بعد",
    }]
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    assert out[0]["notes"] is None


def test_v82_keeps_explicit_notes_label():
    from pipeline.grounding import ground_programs
    doc = """مشروع سقيا الماء
توزيع عبوات الماء
ملاحظات
ينفذ في موسم الحج
"""
    candidate = [{"type": "project", "name": "مشروع سقيا الماء", "notes": "ينفذ في موسم الحج"}]
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    assert out[0]["notes"] == "ينفذ في موسم الحج"


def test_v82_cleans_obvious_arabic_ocr_splits_after_grounding():
    from pipeline.grounding import ground_programs
    doc = """مبادرة ُسلوان
ز يارة المر يضات من ضيفات الرحمن، وتقديم الهدايا لهن
الفئة المستهدفة
المر يضات من ضيفات الرحمن
"""
    candidate = [{
        "type": "project", "name": "مبادرة ُسلوان",
        "description": "ز يارة المر يضات من ضيفات الرحمن، وتقديم الهدايا لهن",
        "target_audience": "المر يضات من ضيفات الرحمن",
    }]
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    assert out[0]["name"] == "مبادرة سلوان"
    assert out[0]["description"].startswith("زيارة المريضات")
    assert out[0]["target_audience"] == "المريضات من ضيفات الرحمن"


def test_v82_supervisor_wrap_goes_to_target_not_delivery_or_notes():
    from pipeline.grounding import ground_programs
    doc = """دورة إدارة استراتيجية لحملة الحج
دورة تدريبية مقدمة لمشرفات الحج
عدد المستفيدين
67
عدد المتطوعين ساعات التطوع مكان التنفيذ الفئة المستهدفة
مشرفات اللجان الثقافية
5 60 عن بعد
في حملات الحج
"""
    candidate = [{
        "type": "project", "name": "دورة إدارة استراتيجية لحملة الحج",
        "beneficiaries_count": 67, "delivery_method": "عن بعد",
        "target_audience": "مشرفات اللجان الثقافية", "notes": "في حملات الحج",
    }]
    out, _ = ground_programs(pl.clean_programs(candidate), doc)
    assert out[0]["target_audience"] == "مشرفات اللجان الثقافية في حملات الحج"
    assert out[0]["delivery_method"] == "عن بعد"
    assert out[0]["notes"] is None


def test_v83_cleans_ocr_splits_even_with_attached_arabic_clitics():
    from pipeline.grounding import _clean_display_text
    samples = {
        "تأهيل وتدر يب فرق تطوعية": "تأهيل وتدريب فرق تطوعية",
        "لتعز يز أثر التجربة": "لتعزيز أثر التجربة",
        "الحجاج والمعتمر ين والزوار": "الحجاج والمعتمرين والزوار",
        "باللغتين العربية والإنجليز ية": "باللغتين العربية والإنجليزية",
        "القرآن الكر يم": "القرآن الكريم",
        "حجاج الداخل والخار ج من العرب": "حجاج الداخل والخارج من العرب",
        "وتوز يع حلوى العيد": "وتوزيع حلوى العيد",
    }
    for raw, expected in samples.items():
        assert _clean_display_text(raw) == expected


def test_v83_cleans_common_parenthesis_and_sentence_spacing_ocr():
    from pipeline.grounding import _clean_display_text
    assert _clean_display_text("تتضمن(:أرًزا مع الدجاج)") == "تتضمن: (أرزًا مع الدجاج)"
    assert _clean_display_text("شملت(:سجادة صلاة)") == "شملت: (سجادة صلاة)"
    assert _clean_display_text("لخدمة الحجاج والمعتمر ين.رّكّز البرنامج") == "لخدمة الحجاج والمعتمرين. رّكّز البرنامج"
    assert _clean_display_text("تقديم الهدايا لهن؛ دعًما نفسيًا") == "تقديم الهدايا لهن؛ دعمًا نفسيًا"


# ------------------------------------------------------------------ v8.4 generalized structure parsing

def test_v84_program_list_context_overrides_name_prefix_and_preserves_source_name():
    from pipeline.grounding import ground_programs
    doc = """التقرير السنوي 2024
أبرز البرامج
هي مسابقة قرآنية لتحفيز الطلاب واكتشاف المواهب :مسابقة رتل
تأهيل وإعداد طلاب المرحلة الثانوية للإمامة في المساجد :محاريب
برنامج حفاظ
حلقات تعتني بالطلاب الحفاظ وتأهيلهم للإجازة بالقرآن.
برنامج يهدف إلى تحسين التلاوة من خلال القراءة التطبيقية للدارس :تصحيح التلاوة للكبار
"""
    # Deliberately give the model a wrong type/prefixed name for one item.
    candidates = pl.clean_programs([
        {"name": "مسابقة رتل", "type": "project"},
        {"name": "برنامج تصحيح التلاوة للكبار", "type": "project"},
    ])
    out, _ = ground_programs(candidates, doc)
    assert [x["name"] for x in out] == [
        "مسابقة رتل", "محاريب", "برنامج حفاظ", "تصحيح التلاوة للكبار"
    ]
    assert all(x["type"] == "program" for x in out)
    assert out[-1]["name"] == "تصحيح التلاوة للكبار"  # no invented "برنامج" prefix
    assert out[0]["description"].startswith("هي مسابقة قرآنية")
    assert out[1]["description"].startswith("تأهيل وإعداد")


def test_v84_project_list_context_classifies_prefixless_names_as_projects():
    from pipeline.grounding import ground_programs
    doc = """Annual Report 2025
أبرز المشاريع
تمكين الأسر: مشروع لدعم الأسر المنتجة وتطوير دخلها
بناء القدرات
مشروع تدريبي لتطوير مهارات العاملين في القطاع.
"""
    out, _ = ground_programs([], doc)
    by_name = {x["name"]: x for x in out}
    assert by_name["تمكين الأسر"]["type"] == "project"
    assert by_name["بناء القدرات"]["type"] == "project"
    assert by_name["تمكين الأسر"]["description"].startswith("مشروع لدعم الأسر")


def test_v84_handles_ltr_and_rtl_colon_layouts_in_same_program_list():
    from pipeline.grounding import ground_programs
    doc = """التقرير السنوي 2025
أبرز البرامج
برنامج ألف: يقدم تدريبًا مهنيًا للشباب.
يقدم جلسات إرشادية للأسر وورش عمل شهرية :برنامج باء
"""
    out, _ = ground_programs([], doc)
    assert [x["name"] for x in out] == ["برنامج ألف", "برنامج باء"]
    assert all(x["type"] == "program" for x in out)
    assert "تدريبًا مهنيًا" in out[0]["description"]
    assert "جلسات إرشادية" in out[1]["description"]


def test_v84_english_program_list_and_unknown_model_shape_fallback():
    from pipeline.grounding import ground_programs
    doc = """Annual Report 2025
Featured Programs
Youth Skills: Provides practical training for young adults.
Family Support
Provides counseling and referral support for families.

Special Services
Community Desk
A staffed help desk for beneficiaries.
"""
    candidates = pl.clean_programs([
        {"name": "Community Desk", "type": "program", "description": "A staffed help desk for beneficiaries."}
    ])
    out, _ = ground_programs(candidates, doc)
    by_name = {x["name"]: x for x in out}
    assert by_name["Youth Skills"]["type"] == "program"
    assert by_name["Family Support"]["type"] == "program"
    # Future/unknown layout not covered by deterministic rules is retained when
    # the model found it and strict grounding verifies the exact source text.
    assert by_name["Community Desk"]["type"] == "program"


def test_v84_does_not_promote_achievement_mentions_or_page_footers():
    from pipeline.grounding import ground_programs
    doc = """البرامج التعليمية
برامج البنين
أبرز البرامج
برنامج المهرة
برنامج يعنى بإعداد الطالب للمسابقات القرآنية وتأهيله للمنافسات المحلية والدولية.
منجزات برنامج المهرة
1. الطالب الأول
المركز الثالث
مسابقة نادي المدينة للقرآن الكريم
2. الطالب الثاني
المركز الأول
مسابقة وزير الرياضة للقرآن الكريم
30
2024 الجمعية الخيرية | التقرير السنوي
صفحة PDF رقم 33
الموارد البشرية:
"""
    out, _ = ground_programs([], doc)
    assert [x["name"] for x in out] == ["برامج البنين", "برنامج المهرة"]
    assert all("مسابقة" not in x["name"] for x in out)
    assert all("صفحة PDF" not in (x.get("description") or "") for x in out)


def test_v84_dedupes_ocr_spacing_and_wrapped_parenthetical_titles():
    from pipeline.grounding import ground_programs
    doc = """15 أولا: برامج إ كرام الزائرات
أولًا: برامج إكرام الزائرات
وصف البرنامج.
28 3.مبادرة ما تيسر منه (تصحيح تلاوة سورة الفاتحة)
مبادرة ما تيسر منه
(تصحيح تلاوة سورة الفاتحة)
تصحيح تلاوة سورة الفاتحة للمستفيدات.
"""
    out, _ = ground_programs([], doc)
    assert [x["name"] for x in out if x["type"] == "program"] == ["برامج إكرام الزائرات"]
    projects = [x for x in out if x["type"] == "project"]
    assert len(projects) == 1
    assert projects[0]["name"] == "مبادرة ما تيسر منه (تصحيح تلاوة سورة الفاتحة)"


def test_v84_report_year_can_be_dominant_even_with_historical_comparison_years():
    from pipeline.grounding import ground_programs
    doc = """التقرير السنوي 2024
أبرز البرامج
خدمة المستفيدين: برنامج يقدم الدعم للمستفيدين.
مقارنة الأداء مع 2023 وتحضير خطة 2025.
2024 الجمعية | التقرير السنوي
2024 الجمعية | التقرير السنوي
"""
    out, _ = ground_programs([], doc)
    assert out[0]["year"] == 2024

# ------------------------------------------------------------------ v8.5 parallel-layout / prose-guard fixes

def test_v85_docx_large_horizontal_gaps_become_column_separators():
    from pipeline.extractors import _clean_docx_paragraph_text
    raw = "إكرام الزائرات                                           إرشاد الزائرات"
    assert _clean_docx_paragraph_text(raw) == "إكرام الزائرات | إرشاد الزائرات"


def test_v85_parallel_program_row_recovers_both_titles_and_drops_prose_names():
    from pipeline.grounding import ground_programs
    doc = """جمعية زائرات البيت الحرام
1445 - 2023
برامجنا
إكرام الزائرات | إرشاد الزائرات
توجيه وإرشاد الزائرات بمبادرات توعوية على عدة مستويات : شرعية أو صحية أو أمنية
رعاية الزائرات
ومنها : زيارة المرضى ومساعدتهم، والعناية بكبيرات السن ورعاية أطفال الزائرات.
تأهيل متطوعات
تأهيل وتدريب فرق تطوعية بلغات مختلفة لخدمة زائرات البيت الحرام، وتلبية حاجاتهم،
وإرشادهم؛ لإتمام نسكهم على الوجه
الصحيح.
نسعى لإكرام زائرات البيت الحرام، وذلك بتقديم وجبات الإطعام والإفطار والسقيا.
الحفاوةبالزائرات
قال رسول الله صلى الله عليه وسلم حديثًا في الحفاوة بالضيف.
أخرجه المصدر في السلسلة الصحيحة
ولذلك نحن نسعى في جمعية زائرات البيت الحرام باستقبال وتوديع وفد الله وتقديم الهدايا لهم.
إنجازاتنا
"""
    candidates = pl.clean_programs([
        {"name": "توجيه وإرشاد الزائرات بمبادرات توعوية على عدة مستويات", "type": "program"},
        {"name": "ومنها", "type": "program"},
        {"name": "أخرجه المصدر في السلسلة الصحيحة", "type": "program"},
    ])
    out, _ = ground_programs(candidates, doc)
    assert [x["name"] for x in out] == [
        "إكرام الزائرات", "إرشاد الزائرات", "رعاية الزائرات",
        "تأهيل متطوعات", "الحفاوة بالزائرات",
    ]
    assert all(x["type"] == "program" for x in out)
    assert all(x["year"] == 2023 for x in out)
    assert "شرعية أو صحية أو أمنية" in out[1]["description"]
    assert out[2]["description"].startswith("ومنها:")
    assert "نسعى لإكرام" in out[0]["description"]
    assert "نسعى في جمعية" in out[-1]["description"]


def test_v85_cover_year_beats_year_like_iban_groups():
    from pipeline.grounding import _unique_document_year
    doc = """جمعية خيرية
1445 - 2023
برامجنا
خدمة المستفيدين
حساباتنا البنكية
SA528 00002 01608 01666 8003
SA520 50000 68205 07382 9000
"""
    assert _unique_document_year(doc) == 2023


def test_v85_rejects_full_description_copied_into_semantic_fields():
    from pipeline.grounding import ground_programs
    desc = "تأهيل وتدريب فرق تطوعية بلغات مختلفة لخدمة الزائرات وتلبية حاجاتهم وإرشادهم"
    doc = f"""التقرير السنوي 2025
برامجنا
تأهيل متطوعات
{desc}
"""
    candidates = pl.clean_programs([{
        "name": "تأهيل متطوعات", "type": "program", "description": desc,
        "beneficiary_value": desc, "delivery_method": desc,
    }])
    out, _ = ground_programs(candidates, doc)
    assert out[0]["beneficiary_value"] is None
    assert out[0]["delivery_method"] is None
